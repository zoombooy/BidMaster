"""解析完整性审计：源 docx 的每个字/表/标题 vs 解析产物逐项比对。"""
import sys, zipfile, re
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, '.')

DOCX = r'C:/Users/张博艺/Desktop/标书文件/extracted/包1/主招标文件/（第2-6章）2026特高压工程服务公开招标采购文件-四次服务.docx'

# 1) 源文档全量文本（XML 级，含文本框）+ 按部件分类
from docx import Document
from docx.table import Table as T
from docx.text.paragraph import Paragraph as P
import lxml.etree as etree

NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
doc = Document(DOCX)

body_items = []          # (kind, obj) 按文档流
for child in doc.element.body.iterchildren():
    tag = child.tag.rsplit('}', 1)[-1]
    if tag == 'p':
        t = P(child, doc).text
        body_items.append(('p', t))
    elif tag == 'tbl':
        body_items.append(('tbl', T(child, doc)))

# XML 级全量文本（含文本框 w:txbxContent）
z = zipfile.ZipFile(DOCX)
doc_xml = z.read('word/document.xml').decode('utf-8')
all_xml_text = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', doc_xml))
# 文本框文本
tb_texts = re.findall(r'<w:txbxContent>.*?</w:txbxContent>', doc_xml, re.S)
tb_text = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', ''.join(tb_texts)))
# 页眉页脚
hf_text = ''
for n in z.namelist():
    if 'header' in n or 'footer' in n:
        if n.endswith('.xml'):
            hf_text += ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', z.read(n).decode('utf-8')))

src_paras = [t for kind, t in body_items if kind == 'p' and t.strip()]
src_tables = [t for kind, t in body_items if kind == 'tbl']
src_table_cells = sum(len(t.rows) * len(t.columns) for t in src_tables)

print(f'源文档: 段落 {len(src_paras)} | 表格 {len(src_tables)} 张 / {src_table_cells} 格 | '
      f'XML全文字符 {len(all_xml_text)} | 文本框 {len(tb_texts)} 处/{len(tb_text)} 字 | 页眉页脚 {len(hf_text)} 字')
print()

# 2) 我们的解析产物
from bidmaster.parser.docx_parser import parse_docx
from bidmaster.ingestion.ledger import ProcessingLedger
parsed = parse_docx(DOCX, 'audit', ProcessingLedger())
ft = parsed.full_text

print(f'解析产物: blocks {len(parsed.blocks)} | tables {len(parsed.tables)} | full_text {len(ft)} 字符')
print()

# 3) 逐段覆盖检查：源段落（>10字）是否出现在 full_text 中
miss = []
for t in src_paras:
    key = t.strip()[:20]
    if len(t.strip()) > 10 and key not in ft:
        miss.append(t)
print(f'段落覆盖: {len(src_paras) - len(miss)}/{len(src_paras)}，'
      f'未覆盖 {len(miss)} 段（{len(miss)/max(len(src_paras),1)*100:.1f}%）')
for t in miss[:8]:
    print('   漏:', t[:70])
print()

# 4) 文本框内容是否被解析
if tb_text.strip():
    probe = tb_text.strip()[:20]
    print(f'文本框: {"已包含" if probe in ft else "未包含 ✗"}（样例: {tb_text[:60]}）')
print()
# 5) 页眉页脚是否被解析
if hf_text.strip():
    probe = hf_text.strip()[:20]
    print(f'页眉页脚: {"已包含" if probe in ft else "未包含 ✗"}（样例: {hf_text[:60]}）')
print()
# 6) 表格对照：源表 vs 解析表（行列规模）
print(f'表格: 源 {len(src_tables)} 张 vs 解析 {len(parsed.tables)} 张')
src_sizes = [(len(t.rows), len(t.columns)) for t in src_tables]
dst_sizes = [(len(r), max(len(x) for x in r) if r else 0) for t in parsed.tables for r in [t.rows]]
print('  源表规模:', src_sizes[:12])
print('  解析表规模:', dst_sizes[:12])
