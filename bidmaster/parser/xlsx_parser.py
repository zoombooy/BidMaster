"""xlsx 解析链路（如"应提交保证金一览表"）：openpyxl → 表格模型。

每个 sheet 产出一个 TableModel + 一个 sheet 名标题块；单元格统一转字符串
（金额保留数值原样，日期转 ISO）。
"""
from __future__ import annotations

from bidmaster.ingestion.ledger import ProcessingLedger
from bidmaster.schemas.document import Block, BlockStyle, ParsedDocument, TableModel


def _cell_str(v) -> str:
    import datetime as _dt
    if v is None:
        return ""
    if isinstance(v, _dt.datetime | _dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx(path: str, doc_id: str, ledger: ProcessingLedger | None = None,
               empty_pages: list[int] | None = None) -> ParsedDocument:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    blocks: list[Block] = []
    tables: list[TableModel] = []
    char_cursor = 0
    b_seq = 0
    t_seq = 0

    def add_block(text: str, btype: str, level: int = 0) -> None:
        nonlocal char_cursor, b_seq
        b_seq += 1
        start = char_cursor
        char_cursor += len(text) + 1
        blocks.append(Block(
            block_id=f"b-{b_seq:05d}", type=btype, page_no=0,
            char_start=start, char_end=start + len(text), text=text,
            style=BlockStyle(heading_level=level)))

    for ws in wb.worksheets:
        add_block(f"工作表：{ws.title}", "heading", level=1)
        rows = [row for row in ws.iter_rows(values_only=True)
                if any(c is not None and str(c).strip() for c in row)]
        if not rows:
            continue
        t_seq += 1
        table_id = f"t-{t_seq:04d}"
        str_rows = [["" if c is None else _cell_str(c) for c in row] for row in rows]
        ncols = max(len(r) for r in str_rows)
        for r in str_rows:
            r += [""] * (ncols - len(r))
        tables.append(TableModel(table_id=table_id, page_nos=[0],
                                 rows=str_rows, header_rows=1))
        summary = "\n".join("\t".join(r) for r in str_rows[:200])
        b_seq += 1
        start = char_cursor
        char_cursor += len(summary) + 1
        blocks.append(Block(
            block_id=f"b-{b_seq:05d}", type="table", page_no=0,
            char_start=start, char_end=start + len(summary), text=summary,
            table_ref=table_id))
    wb.close()

    if ledger is not None:
        ledger.register("file:main", "file")
        ledger.done("file:main")
    return ParsedDocument(
        doc_id=doc_id, file_name=path.split("\\")[-1].split("/")[-1],
        file_type="xlsx", blocks=blocks, tables=tables,
        full_text="\n".join(b.text for b in blocks),
        quality={"warnings": ["xlsx 链路无页码信息"]},
    )
